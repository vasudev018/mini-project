from flask import Flask, render_template, request, redirect, session
from openpyxl import Workbook, load_workbook
from reportlab.pdfgen import canvas
import pandas as pd
import matplotlib.pyplot as plt
import os

app = Flask(__name__)

app.secret_key = "gcs_secret_key"

FILE_NAME = "students.xlsx"

# Create Excel File
if not os.path.exists(FILE_NAME):

    wb = Workbook()
    ws = wb.active

    ws.append(["Name", "Class", "Phone", "Email", "Address"])

    wb.save(FILE_NAME)

# Home


@app.route('/')
def home():
    return render_template('index.html')

# About


@app.route('/about')
def about():
    return render_template('about.html')

# Contact


@app.route('/contact')
def contact():
    return render_template('contact.html')

# Admission


@app.route('/admission', methods=['GET', 'POST'])
def admission():

    if request.method == 'POST':

        name = request.form['name']
        student_class = request.form['class']
        phone = request.form['phone']
        email = request.form['email']
        address = request.form['address']

        wb = load_workbook(FILE_NAME)
        ws = wb.active

        ws.append([name, student_class, phone, email, address])

        wb.save(FILE_NAME)

        return render_template('success.html')

    return render_template('admission.html')

# Admin Login


@app.route('/login', methods=['GET', 'POST'])
def login():

    if request.method == 'POST':

        username = request.form['username']
        password = request.form['password']

        if username == "admin" and password == "1234":

            session['admin'] = True

            return redirect('/dashboard')

    return render_template('login.html')

# Dashboard


@app.route('/dashboard')
def dashboard():

    if 'admin' in session:

        return render_template('dashboard.html')

    return redirect('/login')

# Search Student


@app.route('/search', methods=['GET', 'POST'])
def search():

    students = []

    if request.method == 'POST':

        keyword = request.form['keyword']

        df = pd.read_excel(FILE_NAME)

        result = df[df['Name'].astype(str).str.contains(keyword, case=False)]

        students = result.values.tolist()

    return render_template('search.html', students=students)

# Delete Student


@app.route('/delete/<name>')
def delete(name):

    df = pd.read_excel(FILE_NAME)

    df = df[df['Name'] != name]

    df.to_excel(FILE_NAME, index=False)

    return redirect('/search')

# PDF Report


@app.route('/report')
def report():

    df = pd.read_excel(FILE_NAME)

    if not os.path.exists("reports"):
        os.makedirs("reports")

    pdf = canvas.Canvas("reports/student_report.pdf")

    pdf.drawString(200, 800, "GCS SCHOOL STUDENT REPORT")

    y = 760

    for index, row in df.iterrows():

        text = f"{row['Name']} | {row['Class']} | {row['Phone']}"

        pdf.drawString(50, y, text)

        y -= 20

    pdf.save()

    return "PDF Report Generated Successfully"

# Chart


@app.route('/chart')
def chart():

    df = pd.read_excel(FILE_NAME)

    class_count = df['Class'].value_counts()

    plt.figure(figsize=(6, 5))

    class_count.plot(kind='bar')

    plt.title("Students Per Class")

    plt.savefig('static/images/chart.png')

    return render_template('chart.html')

# Logout


@app.route('/logout')
def logout():

    session.pop('admin', None)

    return redirect('/')


if __name__ == '__main__':
    app.run(debug=True)
